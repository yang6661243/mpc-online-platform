"""MILP optimal solver CLI — python -m microgrid.solver --config solver.yaml"""
import argparse, os, sys, csv, time
from datetime import datetime, timedelta
import yaml
import numpy as np
import pandas as pd
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
from models.benchmark.solver import BenchmarkConfig, solve_benchmark
from microgrid.vpp import build_vpp_window_mask, build_workday_average_baseline
from config_profiles import load_profiled_config

DT = 0.25


def _resolve_path(path):
    """Resolve path: absolute or relative to project root."""
    if os.path.isabs(path):
        return path
    return os.path.join(ROOT, path)


def _read_xlsx_sheet_col(filepath, sheet, keywords):
    """Read a single numeric column from a specific sheet in xlsx, matched by header keyword."""
    wb = load_workbook(filepath, read_only=True, data_only=True)
    ws = wb[sheet]
    headers = [str(c.value or '') for c in next(ws.iter_rows(min_row=1, max_row=1))]
    target = None
    for kw in keywords:
        for idx, h in enumerate(headers):
            if kw.lower() in h.lower():
                target = idx
                break
        if target is not None:
            break
    if target is None:
        raise ValueError(f"None of {keywords} found in headers: {headers}")
    values = []
    for row in ws.iter_rows(min_row=2, min_col=target + 1, max_col=target + 1):
        if row[0].value is not None:
            try:
                values.append(float(row[0].value))
            except (ValueError, TypeError):
                pass
    wb.close()
    return values


def _read_xlsx_sheet_timestamps(filepath, sheet):
    """Read the first column as pandas Timestamps from a specific sheet."""
    wb = load_workbook(filepath, read_only=True, data_only=True)
    ws = wb[sheet]
    timestamps = []
    for row in ws.iter_rows(min_row=2, min_col=1, max_col=1):
        if row[0].value is not None:
            try:
                timestamps.append(pd.Timestamp(row[0].value))
            except (ValueError, TypeError):
                pass
    wb.close()
    return timestamps


def load_config(config_path, profile=None):
    """Load and validate solver YAML config."""
    cfg = load_profiled_config(config_path, profile)
    for section in ['scenario', 'battery', 'device', 'grid', 'cost']:
        if section not in cfg:
            raise ValueError(f"Missing required config section: {section}")
    sc = cfg['scenario']
    for key in ['data_file']:
        if key not in sc:
            raise ValueError(f"Missing required scenario key: {key}")
    return cfg


def write_excel(traj, ts_list, pv_kw, load_kw, buy_prices, cfg, result, output_path):
    """Write solver output Excel with 3 sheets: trajectory, daily, cost summary."""
    T = len(traj)
    DAYS = T / 96
    vpp_cfg = cfg.get('vpp', {})
    vpp_enabled = vpp_cfg.get('enabled', False)
    vpp_charge_price = vpp_cfg.get('charge_price', 0.2)
    vpp_window = [
        vpp_enabled and vpp_cfg.get('start_hour', 11) <= ts_list[i].hour < vpp_cfg.get('end_hour', 13)
        for i in range(T)
    ]
    vpp_baseline = list(vpp_cfg.get('baseline_kw') or vpp_cfg.get('baseline') or [])
    if len(vpp_baseline) < T:
        vpp_baseline.extend([0.0] * (T - len(vpp_baseline)))
    vpp_response = list(result.vpp_response_kw or [])
    if len(vpp_response) < T:
        vpp_response.extend([0.0] * (T - len(vpp_response)))
    vpp_step_benefit = [
        max(buy_prices[i] - vpp_charge_price, 0.0) * vpp_response[i] * DT
        for i in range(T)
    ]
    hf = Font(bold=True)
    hfl = PatternFill("solid", fgColor="DDEEFF")
    ha = Alignment(horizontal="center")
    wb = Workbook()

    def set_h(ws, hdrs):
        for c, h in enumerate(hdrs, 1):
            cell = ws.cell(row=1, column=c, value=h)
            cell.font = hf; cell.fill = hfl; cell.alignment = ha

    # Sheet 1: 15min_trajectory
    ws1 = wb.active
    ws1.title = '15min_trajectory'
    hdr = ['时间', '光伏功率(kW)', '负载功率(kW)', 'SOC', '电池功率(kW)',
           '电网功率(kW)', '购电价(元/kWh)', '售电价(元/kWh)']
    hdr.extend(['互济窗口', '互济基线(kW)', '互济结算功率(kW)', '互济结算电量(kWh)', '互济收益(元)'])
    set_h(ws1, hdr)
    for i in range(T):
        ws1.append([ts_list[i].strftime('%Y-%m-%d %H:%M'),
                    round(pv_kw[i], 4), round(load_kw[i], 4),
                    round(traj[i][0], 4), round(traj[i][1], 4),
                    round(traj[i][2], 4), buy_prices[i], 0.0])
        ws1.cell(row=i + 2, column=9, value='是' if vpp_window[i] else '否')
        ws1.cell(row=i + 2, column=10, value=round(vpp_baseline[i], 4))
        ws1.cell(row=i + 2, column=11, value=round(vpp_response[i], 4))
        ws1.cell(row=i + 2, column=12, value=round(vpp_response[i] * DT, 4))
        ws1.cell(row=i + 2, column=13, value=round(vpp_step_benefit[i], 4))
    for c in range(1, 14):
        ws1.column_dimensions[get_column_letter(c)].width = 20

    # Sheet 2: daily_summary
    ws2 = wb.create_sheet('daily_summary')
    set_h(ws2, ['天数', '日期', '起始SOC', '结束SOC', '最低SOC', '最高SOC',
                '用电量(kWh)', '充电量(kWh)', '放电量(kWh)', '购电费(元)',
                '衰减(元)', '净成本(元)', '互济结算电量(kWh)', '互济收益(元)'])
    for d in range(int(DAYS)):
        s, e = d * 96, (d + 1) * 96
        sd = traj[s:e]
        soc_slice = [r[0] for r in sd]
        bat_slice = [r[1] for r in sd]
        gi_slice = [r[2] for r in sd]
        load_energy = sum(load_kw[s:e]) * DT
        chg_energy = sum(max(0, b) for b in bat_slice) * DT
        dis_energy = sum(max(0, -b) for b in bat_slice) * DT
        purchase = sum(buy_prices[i] * max(0, gi_slice[i - s]) * DT for i in range(s, e))
        vpp_energy = sum(vpp_response[s:e]) * DT
        vpp_benefit = sum(vpp_step_benefit[s:e])
        deg = sum(0.05 * abs(b) * DT for b in bat_slice)
        net = purchase - vpp_benefit + deg
        ws2.append([d + 1, ts_list[s].strftime('%Y-%m-%d'),
                    round(soc_slice[0], 4), round(soc_slice[-1], 4),
                    round(min(soc_slice), 4), round(max(soc_slice), 4),
                    round(load_energy, 1), round(chg_energy, 1), round(dis_energy, 1),
                    round(purchase, 2), round(deg, 2), round(net, 2),
                    round(vpp_energy, 2), round(vpp_benefit, 2)])
    for c in range(1, 15):
        ws2.column_dimensions[get_column_letter(c)].width = 16

    # Sheet 3: cost_summary
    ws3 = wb.create_sheet('cost_summary')
    set_h(ws3, ['指标', '数值(元)'])
    ws3.column_dimensions['A'].width = 22
    ws3.column_dimensions['B'].width = 16

    purchase_cost = sum(buy_prices[i] * max(0, traj[i][2]) * DT for i in range(T))
    vpp_benefit = result.vpp_benefit
    vpp_energy_total = sum(vpp_response) * DT
    vpp_avg_price = (
        vpp_charge_price if vpp_energy_total > 0 else 0.0
    )
    degradation = sum(0.05 * abs(traj[i][1]) * DT for i in range(T))
    peak = max(traj[i][2] for i in range(T))
    demand_charge = cfg['cost']['demand_rate'] * peak * DAYS / 30
    comprehensive = purchase_cost - vpp_benefit + degradation + demand_charge

    total_pv = sum(pv_kw) * DT
    pv_curtailed = sum(max(0, pv_kw[i] - load_kw[i] - max(0, traj[i][1])) for i in range(T)) * DT
    renewable_rate = (1 - pv_curtailed / total_pv) if total_pv > 0 else 1.0

    rows = [
        ('购电费', round(purchase_cost, 2)),
        ('售电收益', 0.0),
        ('互济总收益(元)', round(vpp_benefit, 2)),
        ('互济总结算电量(kWh)', round(vpp_energy_total, 2)),
        ('等效互济均价(元/kWh)', round(vpp_avg_price, 4)),
        ('电池衰减', round(degradation, 2)),
        ('需量费', round(demand_charge, 2)),
        ('容量电费', 0.0),
        ('综合成本', round(comprehensive, 2)),
        ('峰值需量(kW)', round(peak, 1)),
        ('可再生能源利用率', round(renewable_rate, 4)),
    ]
    for i, (k, v) in enumerate(rows):
        ws3.cell(row=i + 2, column=1, value=k)
        ws3.cell(row=i + 2, column=2, value=v)

    try:
        wb.save(output_path)
    except PermissionError as exc:
        raise PermissionError(
            f"无法写入输出文件：{output_path}。请先关闭已打开的同名 Excel 文件，或修改配置里的 output 路径。"
        ) from exc
    return comprehensive


def main():
    parser = argparse.ArgumentParser(description="MILP optimal solver")
    parser.add_argument("--config", required=True, help="Path to solver YAML config")
    parser.add_argument("--profile", help="Named profile inside a profiled YAML config")
    args = parser.parse_args()

    config_path = _resolve_path(args.config)
    if not os.path.exists(config_path):
        print(f"ERROR: config file not found: {config_path}")
        sys.exit(1)

    cfg = load_config(config_path, args.profile)
    profile_label = f" [profile={cfg.get('_profile')}]" if cfg.get('_profile') else ""
    print(f"Loading config: {config_path}{profile_label}")

    sc = cfg['scenario']; bt = cfg['battery']; dv = cfg['device']
    gd = cfg['grid']; cs = cfg['cost']; sv = cfg.get('solver', {})
    out_path = _resolve_path(cfg.get('output', 'outputs/solver_result.xlsx'))

    # Load scenario data from single Excel with multiple sheets
    data_file = _resolve_path(sc['data_file'])
    sheets_cfg = sc.get('sheets', {})
    load_sheet = sheets_cfg.get('load', 'load')
    pv_sheet = sheets_cfg.get('pv_wind', 'pv')
    price_sheet = sheets_cfg.get('price', 'price')

    print(f"Loading data from: {data_file}")
    load_raw = _read_xlsx_sheet_col(data_file, load_sheet,
                                     ['负荷', 'load', '有功', 'demand'])
    pv_raw = _read_xlsx_sheet_col(data_file, pv_sheet,
                                   ['irradiance', '辐照', 'ghi', 'solar'])
    wind_raw = _read_xlsx_sheet_col(data_file, pv_sheet,
                                     ['wind', '风速', 'wind_speed'])
    buy_raw = _read_xlsx_sheet_col(data_file, price_sheet,
                                    ['buy_price', '购电价'])
    sell_raw = _read_xlsx_sheet_col(data_file, price_sheet,
                                     ['sell_price', '售电价'])
    ts_list = _read_xlsx_sheet_timestamps(data_file, load_sheet)

    T_raw = min(len(load_raw), len(pv_raw), len(wind_raw),
                len(buy_raw), len(sell_raw), len(ts_list))
    days_req = sv.get('days', 30)
    T = min(days_req * 96, T_raw)
    DAYS = T / 96

    # Convert to kW
    load_base = dv.get('load_base_kw', 1000)
    load_kw = [max(0, load_raw[i]) * load_base for i in range(T)]
    pv_kw = [min(1.0, pv_raw[i] / 1000.0) * dv.get('pv_capacity_kw', 0) * dv.get('pv_efficiency', 0)
             for i in range(T)]
    wind_kw = [min(1.0, (wind_raw[i] / 3.6) / 12.0) * dv.get('wind_capacity_kw', 0) * dv.get('wind_efficiency', 0)
               for i in range(T)]
    buy_prices = buy_raw[:T]
    sell_prices = sell_raw[:T]
    ts_list = ts_list[:T]
    uncontrolled_grid_kw = [
        max(0.0, load_kw[i] - pv_kw[i] - wind_kw[i])
        for i in range(T)
    ]
    vpp = cfg.get('vpp', {})
    vpp_enabled = vpp.get('enabled', False)
    vpp_baseline = build_workday_average_baseline(
        uncontrolled_grid_kw,
        baseline_days=vpp.get('baseline_days', 5),
        steps_per_day=vpp.get('steps_per_day', 96),
    ) if vpp_enabled else []
    vpp_window = build_vpp_window_mask(
        ts_list,
        start_hour=vpp.get('start_hour', 11),
        end_hour=vpp.get('end_hour', 13),
    ) if vpp_enabled else []

    print(f"  {T} steps ({DAYS:.0f} days), Load max={max(load_kw):.0f}kW, PV max={max(pv_kw):.0f}kW")

    # Build BenchmarkConfig
    bench_cfg = BenchmarkConfig(
        battery_capacity_kwh=bt['capacity_kwh'],
        battery_charge_max_kw=bt['charge_max_kw'],
        battery_discharge_max_kw=bt['discharge_max_kw'],
        battery_soc_init=bt.get('soc_init', 0.5),
        battery_soc_min=bt.get('soc_min', 0.1),
        battery_soc_max=bt.get('soc_max', 0.9),
        battery_charge_eff=bt.get('charge_eff', 0.95),
        battery_discharge_eff=bt.get('discharge_eff', 0.95),
        grid_import_max_kw=gd.get('import_max_kw', 5000),
        grid_export_max_kw=gd.get('export_max_kw', 0),
        transformer_capacity_kw=gd.get('transformer_capacity_kw', 5000),
        anti_backflow=gd.get('anti_backflow', True),
        c_deg=cs.get('c_deg', 0.05),
        r_demand=cs.get('demand_rate', 0),
        r_capacity=cs.get('capacity_rate', 0),
        use_milp=sv.get('use_milp', True),
        terminal_constraint=sv.get('terminal_constraint', False),
        optimization_billing_days=cs.get('billing_days', 30),
        vpp_enabled=vpp_enabled,
        vpp_charge_price=vpp.get('charge_price', 0.2),
        vpp_baseline_kw=vpp_baseline,
        vpp_window_mask=vpp_window,
        vpp_response_cap_kw=vpp.get('response_cap_kw'),
    )

    # Solve
    print("Solving MILP...")
    t0 = time.time()
    r = solve_benchmark(pv_kw, wind_kw, load_kw, buy_prices, sell_prices, config=bench_cfg, verbose=False)
    elapsed = time.time() - t0

    demand_charge = cs['demand_rate'] * r.peak_demand_kw * DAYS / 30
    basic = r.purchase_cost - r.export_revenue - r.vpp_benefit + r.degradation_cost
    comprehensive = basic + demand_charge
    print(f"  Peak={r.peak_demand_kw:.1f}kW, Purchase={r.purchase_cost:.0f}, "
          f"Deg={r.degradation_cost:.0f}, Demand={demand_charge:.0f}, "
          f"Comprehensive={comprehensive:.0f} [{elapsed:.1f}s]")

    # Build trajectory
    opt_bat = r.battery_power if r.battery_power else [0] * T
    opt_soc = r.soc_trajectory if r.soc_trajectory else [0.5] * T
    opt_grid = r.grid_import if r.grid_import else [0] * T
    traj = [(opt_soc[i], opt_bat[i], opt_grid[i]) for i in range(T)]

    # Write output
    os.makedirs(os.path.dirname(out_path) or '.', exist_ok=True)
    report_cfg = dict(cfg)
    report_vpp = dict(vpp)
    report_vpp['baseline_kw'] = vpp_baseline
    report_cfg['vpp'] = report_vpp
    write_excel(traj, ts_list, pv_kw, load_kw, buy_prices, report_cfg, r, out_path)
    print(f"Done: {out_path}")

    # Optional: compare with actual
    actual_cfg = cfg.get('actual')
    if actual_cfg and actual_cfg.get('file'):
        print("\n--- Comparing with factory actual ---")
        actual_file = _resolve_path(actual_cfg['file'])
        actual_sheet = actual_cfg.get('sheet', 'Sheet0')
        df_act = pd.read_excel(actual_file, sheet_name=actual_sheet)

        # Read actual data using keyword matching
        # Load the full raw columns for grid and battery
        wb_act = load_workbook(actual_file, read_only=True, data_only=True)
        ws_act = wb_act[actual_sheet]
        headers_act = [str(c.value or '') for c in next(ws_act.iter_rows(min_row=1, max_row=1))]
        grid_col_idx = None; bat_col_idx = None
        for idx, h in enumerate(headers_act):
            hl = h.lower()
            if grid_col_idx is None and ('电网' in hl or 'grid' in hl):
                grid_col_idx = idx
            if bat_col_idx is None and ('储能' in hl or 'battery' in hl or 'bat' in hl):
                bat_col_idx = idx
        actual_grid_vals = []
        actual_bat_vals = []
        for row in ws_act.iter_rows(min_row=2, min_col=1, max_col=max(grid_col_idx or 0, bat_col_idx or 0) + 1):
            if grid_col_idx is not None and grid_col_idx < len(row) and row[grid_col_idx].value is not None:
                try: actual_grid_vals.append(float(row[grid_col_idx].value))
                except (ValueError, TypeError): actual_grid_vals.append(0.0)
            if bat_col_idx is not None and bat_col_idx < len(row) and row[bat_col_idx].value is not None:
                try: actual_bat_vals.append(float(row[bat_col_idx].value))
                except (ValueError, TypeError): actual_bat_vals.append(0.0)
        wb_act.close()

        if actual_grid_vals and actual_bat_vals:
            actual_grid = actual_grid_vals[:T]
            actual_bat = actual_bat_vals[:T]

            act_purchase = sum(buy_prices[i] * max(0, actual_grid[i]) * DT for i in range(T))
            act_deg = sum(0.05 * abs(actual_bat[i]) * DT for i in range(T))
            act_peak = max(actual_grid)
            act_demand = cs['demand_rate'] * act_peak * DAYS / 30
            act_ctrl = act_purchase + act_deg + act_demand

            opt_ctrl = comprehensive
            achieve = (opt_ctrl / act_ctrl * 100) if act_ctrl > 0 else 0

            print(f"  Factory: purchase={act_purchase:.0f} deg={act_deg:.0f} "
                  f"demand={act_demand:.0f} peak={act_peak:.1f} total={act_ctrl:.0f}")
            print(f"  Optimal: purchase={r.purchase_cost:.0f} deg={r.degradation_cost:.0f} "
                  f"demand={demand_charge:.0f} peak={r.peak_demand_kw:.1f} total={opt_ctrl:.0f}")
            print(f"  Achievement rate: {achieve:.1f}%")
        else:
            print("  WARNING: Could not auto-detect grid/battery columns in actual file")


if __name__ == '__main__':
    main()
