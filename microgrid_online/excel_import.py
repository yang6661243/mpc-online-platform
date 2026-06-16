from __future__ import annotations

import uuid
from datetime import datetime
from io import BytesIO
from typing import Sequence

import openpyxl
from sqlalchemy import select
from sqlalchemy.orm import Session

from microgrid_online.models import MpcRun, StrategyComparison, StrategyCurvePoint, utc_now

COLUMN_MAP = {
    "时间": "time",
    "电网功率(kW)": "mpc_grid_power_kw",
    "电池功率(kW)": "mpc_battery_power_kw",
    "SOC": "mpc_soc",
    "负荷功率(kW)": "mpc_load_kw",
    "光伏出力(kW)": "mpc_pv_kw",
    "购电价(元/kWh)": "buy_price",
    "售电价(元/kWh)": "sell_price",
}

DATE_FORMATS = [
    "%Y-%m-%d %H:%M:%S",
    "%Y/%m/%d %H:%M:%S",
    "%Y-%m-%dT%H:%M:%S",
]


def _parse_time(value: object) -> datetime:
    if isinstance(value, datetime):
        return value
    if isinstance(value, str):
        for fmt in DATE_FORMATS:
            try:
                return datetime.strptime(value.strip(), fmt)
            except ValueError:
                continue
    raise ValueError(f"无法解析时间: {value!r}")


def _read_sheet(file_bytes: bytes, sheet_name: str = "15min_trajectory") -> list[dict]:
    wb = openpyxl.load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        wb.close()
        raise ValueError(f"Excel 中未找到 '{sheet_name}' sheet，可用: {wb.sheetnames}")

    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if len(rows) < 2:
        raise ValueError("Excel 至少需要一行表头 + 一行数据")

    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    data = []
    for row in rows[1:]:
        if all(v is None for v in row):
            continue
        record: dict = {}
        for idx, header in enumerate(headers):
            if header in COLUMN_MAP:
                record[COLUMN_MAP[header]] = row[idx] if idx < len(row) else None
        if "time" not in record:
            continue
        record["time"] = _parse_time(record["time"])
        # 净负荷 = 负荷 - 光伏
        load = record.get("mpc_load_kw")
        pv = record.get("mpc_pv_kw")
        if load is not None and pv is not None:
            record["load_minus_pv_kw"] = round(float(load) - float(pv), 6)
        else:
            record["load_minus_pv_kw"] = None
        # Convert floats
        for key in record:
            if key != "time" and record[key] is not None:
                try:
                    record[key] = float(record[key])
                except (ValueError, TypeError):
                    record[key] = None
        data.append(record)
    return data


def _read_cost_summary(file_bytes: bytes) -> dict:
    try:
        wb = openpyxl.load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
        if "cost_summary" not in wb.sheetnames:
            wb.close()
            return {}
        ws = wb["cost_summary"]
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
        result: dict = {}
        for row in rows[1:]:
            if row[0] is None:
                continue
            key = str(row[0]).strip()
            val = row[1]
            if isinstance(val, str) and "%" in val:
                val = float(val.replace("%", "")) / 100
            try:
                val = float(val)
            except (ValueError, TypeError):
                pass
            result[key] = val
        return result
    except Exception:
        return {}


def _compute_peak_kw(points: Sequence[dict], field: str) -> float | None:
    values = [p[field] for p in points if p.get(field) is not None]
    if not values:
        return None
    return round(max(values), 4)


def _detect_month_label(first_time: datetime) -> str:
    return f"{first_time.month}月"


def import_mpc_run_from_excel(
    session: Session,
    *,
    plant_id: str,
    file_bytes: bytes,
    profile: str,
) -> dict:
    rows = _read_sheet(file_bytes)
    if not rows:
        raise ValueError("Excel 中无有效数据行")

    first_time = rows[0]["time"]
    month_label = _detect_month_label(first_time)

    # 自动拼上月前缀：避免重复拼接
    if not profile.startswith(f"{month_label}-"):
        full_profile = f"{month_label}-{profile}"
    else:
        full_profile = profile

    run_id = f"{plant_id}_{full_profile}_{uuid.uuid4().hex[:8]}"

    # 检查 run_id 是否已存在
    existing = session.scalar(select(MpcRun).where(MpcRun.run_id == run_id))
    if existing is not None:
        raise ValueError(f"run_id 已存在: {run_id}")

    last_time = rows[-1]["time"]

    # 创建 MpcRun
    run = MpcRun(
        run_id=run_id,
        plant_id=plant_id,
        profile=full_profile,
        status="imported",
        input_start_time=first_time,
        input_end_time=last_time,
        started_at=utc_now(),
        finished_at=utc_now(),
    )
    session.add(run)

    # 插入曲线点
    # 工厂策略 = 无电池调度: 电网=负荷-光伏, 电池=0, SOC=固定初始值
    # MPC 策略 = 优化后: 电网/电池/SOC 来自 Excel MPC 结果
    curve_points = []
    for row in rows:
        load = row.get("mpc_load_kw")
        pv = row.get("mpc_pv_kw")
        factory_grid = round(load - pv, 6) if load is not None and pv is not None else None

        point = StrategyCurvePoint(
            run_id=run_id,
            plant_id=plant_id,
            time=row["time"],
            # 工厂策略（无电池）
            actual_grid_power_kw=factory_grid,
            actual_battery_power_kw=0.0,
            actual_soc=0.5,
            actual_load_kw=load,
            actual_pv_kw=pv,
            # MPC 策略（优化后）
            load_minus_pv_kw=row.get("load_minus_pv_kw"),
            mpc_grid_power_kw=row.get("mpc_grid_power_kw"),
            mpc_battery_power_kw=row.get("mpc_battery_power_kw"),
            mpc_soc=row.get("mpc_soc"),
            mpc_load_kw=load,
            mpc_pv_kw=pv,
            buy_price=row.get("buy_price"),
            sell_price=row.get("sell_price"),
        )
        session.add(point)
        curve_points.append(point)

    # 读取 cost_summary 用于 StrategyComparison
    costs = _read_cost_summary(file_bytes)

    mpc_peak_kw = _compute_peak_kw(rows, "mpc_grid_power_kw")
    # 工厂峰值 = 无电池调度的电网最大需量
    factory_peak_kw = _compute_peak_kw([{
        "grid": round(r["mpc_load_kw"] - r["mpc_pv_kw"], 6)
        if r.get("mpc_load_kw") is not None and r.get("mpc_pv_kw") is not None else None
    } for r in rows], "grid")

    comparison = StrategyComparison(
        run_id=run_id,
        plant_id=plant_id,
        actual_peak_kw=factory_peak_kw,
        mpc_peak_kw=mpc_peak_kw,
        peak_reduction_kw=round(factory_peak_kw - mpc_peak_kw, 4) if factory_peak_kw is not None and mpc_peak_kw is not None else None,
        peak_reduction_pct=round((factory_peak_kw - mpc_peak_kw) / factory_peak_kw * 100, 2) if factory_peak_kw and factory_peak_kw > 0 and mpc_peak_kw is not None else None,
        actual_cost_yuan=costs.get("MILP 基准"),
        mpc_cost_yuan=costs.get("购电费(元)"),
        cost_saving_yuan=costs.get("套利收益(元)"),
        cost_saving_pct=None,
    )
    session.add(comparison)

    session.commit()

    return {
        "success": True,
        "run_id": run_id,
        "plant_id": plant_id,
        "profile": full_profile,
        "point_count": len(curve_points),
        "time_range": {
            "start": first_time.isoformat() if isinstance(first_time, datetime) else str(first_time),
            "end": last_time.isoformat() if isinstance(last_time, datetime) else str(last_time),
        },
        "mpc_peak_kw": mpc_peak_kw,
    }
