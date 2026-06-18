"""MPC rolling dispatch — python -m mpc.microgrid.mpc --config mpc.yaml"""
import argparse, os, sys, time
from datetime import datetime, timedelta
import yaml
import numpy as np
import pandas as pd
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, ROOT)
from mpc.solvers.benchmark.solver import BenchmarkConfig, solve_benchmark
from mpc.solvers.forecasting import LoadForecaster
from mpc.solvers.forecasting.covariates import build_calendar_covariates, align_covariates
from mpc.microgrid.vpp import (
    build_vpp_window_mask,
    build_workday_average_baseline,
    compute_vpp_step_benefit,
)
from mpc.microgrid.config_profiles import load_profiled_config
from mpc.microgrid.battery_ops import (
    apply_realtime_load_compensation,
    build_dynamic_reserve_soc,
    build_reserve_energy_request,
    estimate_under_forecast_margin,
    _lightgbm_min_history_steps,
    limit_battery_power_by_ramp,
    limit_battery_power_by_soc,
)

DT = 0.25


def _resolve_path(path):
    if os.path.isabs(path):
        return path
    return os.path.join(ROOT, path)


def _read_xlsx_sheet_col(filepath, sheet, keywords):
    wb = load_workbook(filepath, read_only=True, data_only=True)
    ws = wb[sheet]
    headers = [str(c.value or '') for c in next(ws.iter_rows(min_row=1, max_row=1))]
    target = None
    for kw in keywords:
        for idx, h in enumerate(headers):
            if kw.lower() in h.lower():
                target = idx
                break
        if target is not None:
            break
    if target is None:
        raise ValueError(f"None of {keywords} found in headers: {headers}")
    values = []
    for row in ws.iter_rows(min_row=2, min_col=target + 1, max_col=target + 1):
        if row[0].value is not None:
            try:
                values.append(float(row[0].value))
            except (ValueError, TypeError):
                pass
    wb.close()
    return values


def _read_xlsx_sheet_timestamps(filepath, sheet):
    wb = load_workbook(filepath, read_only=True, data_only=True)
    ws = wb[sheet]
    timestamps = []
    for row in ws.iter_rows(min_row=2, min_col=1, max_col=1):
        if row[0].value is not None:
            try:
                timestamps.append(pd.Timestamp(row[0].value))
            except (ValueError, TypeError):
                pass
    wb.close()
    return timestamps


def _target_peak_value_is_set(value):
    if value is None:
        return False
    text = str(value).strip()
    if text == '' or text.lower() in ('null', 'none'):
        return False
    try:
        return float(value) > 0
    except (TypeError, ValueError):
        return True


def _resolve_target_peak_mode(mpc_cfg):
    """Return the explicit target-peak planning mode.

    Backward compatibility:
    - target_peak_kw set without a mode means manual.
    - target_peak_kw null without a mode keeps the old oracle behavior.
    """
    mode = mpc_cfg.get('target_peak_mode')
    if mode is None:
        return 'manual' if _target_peak_value_is_set(mpc_cfg.get('target_peak_kw')) else 'oracle_full_period'

    mode = str(mode).strip().lower()
    aliases = {
        'oracle': 'oracle_full_period',
        'full_period': 'oracle_full_period',
        'forecast': 'forecast_month_plan',
        'deployable': 'forecast_month_plan',
    }
    mode = aliases.get(mode, mode)
    valid = {'manual', 'oracle_full_period', 'forecast_month_plan'}
    if mode not in valid:
        raise ValueError(
            f"Invalid mpc.target_peak_mode={mode!r}. "
            f"Expected one of: {', '.join(sorted(valid))}"
        )
    return mode


def _parse_manual_target_peak(mpc_cfg):
    value = mpc_cfg.get('target_peak_kw')
    if not _target_peak_value_is_set(value):
        raise ValueError("mpc.target_peak_mode=manual requires a positive target_peak_kw.")
    target = float(value)
    if target <= 0:
        raise ValueError("mpc.target_peak_kw must be positive when target_peak_mode=manual.")
    return target


def _compose_forecast_history(load_ratio_full, abs_t, history_prefix=None):
    """Build forecast history without including the current/future simulation step."""
    current_history = np.asarray(load_ratio_full[:abs_t], dtype=np.float64)
    if history_prefix is None or len(history_prefix) == 0:
        return current_history
    return np.concatenate([
        np.asarray(history_prefix, dtype=np.float64),
        current_history,
    ])


def _forecast_history_config(cfg, forecaster=None):
    hist_cfg = cfg.get('forecast_history') or {}
    if hist_cfg:
        return hist_cfg, 'forecast_history'

    fc_cfg = getattr(forecaster, '_cfg', {}) or {}
    data_cfg = fc_cfg.get('data', {}) or {}
    data_file = data_cfg.get('file') or data_cfg.get('scenario_file')
    data_sheet = data_cfg.get('sheet')
    if data_file and data_sheet:
        return {'data_file': data_file, 'sheet': data_sheet}, 'forecaster.data'

    return {}, ''


def _load_forecast_history_prefix(cfg, forecaster, load_base_kw):
    """Load historical load ratios used to warm-start deployable forecasts."""
    hist_cfg, source = _forecast_history_config(cfg, forecaster)
    if not hist_cfg:
        return np.array([], dtype=np.float64), ''

    data_file = hist_cfg.get('data_file') or hist_cfg.get('file')
    sheet = hist_cfg.get('sheet') or hist_cfg.get('load_sheet')
    if not data_file or not sheet:
        raise ValueError(
            "forecast_history requires data_file/file and sheet/load_sheet "
            "when LightGBM needs historical warm-start data."
        )

    values = _read_xlsx_sheet_col(
        _resolve_path(data_file),
        sheet,
        ['负荷', 'load', '有功', 'demand', 'value'],
    )
    unit = str(hist_cfg.get('unit', hist_cfg.get('value_unit', 'ratio'))).strip().lower()
    if unit == 'kw':
        if load_base_kw <= 0:
            raise ValueError("device.load_base_kw must be positive when forecast_history.unit=kw.")
        ratios = [float(v) / load_base_kw for v in values]
    elif unit in ('ratio', 'pu', 'p.u.', 'per_unit'):
        ratios = [float(v) for v in values]
    else:
        raise ValueError("forecast_history.unit must be 'ratio' or 'kw'.")

    ratios = [max(0.0, min(1.0, r)) for r in ratios]
    label = f"{source}:{data_file}[{sheet}]"
    return np.asarray(ratios, dtype=np.float64), label


def _load_target_peak_forecast_inputs(plan_cfg, default_sheets, load_base_kw, device_cfg, fallback_price):
    """Read a deployable month-plan forecast workbook for target-peak planning."""
    data_file = plan_cfg.get('data_file') or plan_cfg.get('file')
    if not data_file:
        raise ValueError(
            "mpc.target_peak_mode=forecast_month_plan requires "
            "mpc.target_peak_forecast.data_file."
        )

    sheets_cfg = plan_cfg.get('sheets', {})
    load_sheet = sheets_cfg.get('load') or plan_cfg.get('load_sheet') or default_sheets.get('load', 'load')
    pv_sheet = sheets_cfg.get('pv_wind') or plan_cfg.get('pv_wind_sheet') or default_sheets.get('pv_wind', 'pv')
    price_sheet = sheets_cfg.get('price') or plan_cfg.get('price_sheet')

    path = _resolve_path(data_file)
    load_raw = _read_xlsx_sheet_col(path, load_sheet, ['负荷', 'load', '有功', 'demand', 'value'])
    pv_raw = _read_xlsx_sheet_col(path, pv_sheet, ['irradiance', '辐照', 'ghi', 'solar'])
    wind_raw = _read_xlsx_sheet_col(path, pv_sheet, ['wind', '风速', 'wind_speed'])

    load_unit = str(plan_cfg.get('load_unit', 'ratio')).strip().lower()
    if load_unit == 'kw':
        load_kw = [max(0.0, float(v)) for v in load_raw]
    elif load_unit in ('ratio', 'pu', 'p.u.', 'per_unit'):
        load_kw = [max(0.0, min(1.0, float(v))) * load_base_kw for v in load_raw]
    else:
        raise ValueError("mpc.target_peak_forecast.load_unit must be 'ratio' or 'kw'.")

    pv_cap = device_cfg.get('pv_capacity_kw', 0)
    pv_eff = device_cfg.get('pv_efficiency', 0)
    wind_cap = device_cfg.get('wind_capacity_kw', 0)
    wind_eff = device_cfg.get('wind_efficiency', 0)
    pv_kw = [min(1.0, float(p) / 1000.0) * pv_cap * pv_eff for p in pv_raw]
    wind_kw = [min(1.0, (float(w) / 3.6) / 12.0) * wind_cap * wind_eff for w in wind_raw]

    if price_sheet:
        buy = _read_xlsx_sheet_col(path, price_sheet, ['buy_price', '购电价'])
        sell = _read_xlsx_sheet_col(path, price_sheet, ['sell_price', '售电价'])
    else:
        buy, sell = fallback_price

    return pv_kw, wind_kw, load_kw, buy, sell


def main():
    parser = argparse.ArgumentParser(description="MPC rolling dispatch")
    parser.add_argument("--config", required=True, help="Path to MPC YAML config")
    parser.add_argument("--profile", help="Named profile inside a profiled YAML config")
    args = parser.parse_args()

    config_path = _resolve_path(args.config)
    if not os.path.exists(config_path):
        print(f"ERROR: config file not found: {config_path}")
        sys.exit(1)

    cfg = load_profiled_config(config_path, args.profile)
    profile_label = f" [profile={cfg.get('_profile')}]" if cfg.get('_profile') else ""
    print(f"Loading config: {config_path}{profile_label}")

    for section in ['scenario', 'battery', 'device', 'grid', 'cost', 'mpc']:
        if section not in cfg:
            raise ValueError(f"Missing required config section: {section}")

    scenario = cfg['scenario']; battery = cfg['battery']; device = cfg['device']
    grid = cfg['grid']; cost = cfg['cost']; mpc_cfg = cfg['mpc']
    out_path = _resolve_path(cfg.get('output', 'outputs/mpc_result.xlsx'))

    # Load scenario data from single Excel with multiple sheets
    data_file = _resolve_path(scenario['data_file'])
    sheets_cfg = scenario.get('sheets', {})
    load_sheet = sheets_cfg.get('load', 'load')
    pv_sheet = sheets_cfg.get('pv_wind', 'pv')
    price_sheet = sheets_cfg.get('price', 'price')

    print(f"Loading data from: {data_file}")
    load_raw = _read_xlsx_sheet_col(data_file, load_sheet,
                                     ['负荷', 'load', '有功', 'demand'])
    pv_raw = _read_xlsx_sheet_col(data_file, pv_sheet,
                                   ['irradiance', '辐照', 'ghi', 'solar'])
    wind_raw = _read_xlsx_sheet_col(data_file, pv_sheet,
                                     ['wind', '风速', 'wind_speed'])
    buy_raw = _read_xlsx_sheet_col(data_file, price_sheet,
                                    ['buy_price', '购电价'])
    sell_raw = _read_xlsx_sheet_col(data_file, price_sheet,
                                     ['sell_price', '售电价'])
    ts_all = _read_xlsx_sheet_timestamps(data_file, load_sheet)

    total_steps = min(len(load_raw), len(pv_raw), len(wind_raw),
                len(buy_raw), len(sell_raw), len(ts_all))
    start_step = mpc_cfg.get('start_step', 0)
    run_days = mpc_cfg.get('days', 30)
    run_steps = min(run_days * 96, total_steps - start_step)
    horizon = mpc_cfg.get('horizon_steps', 192)
    start_time = ts_all[start_step]  # Correct start time from actual data

    load_base = device.get('load_base_kw', 1000)
    pv_cap = device.get('pv_capacity_kw', 0)
    pv_eff = device.get('pv_efficiency', 0)
    wind_cap = device.get('wind_capacity_kw', 0)
    wind_eff = device.get('wind_efficiency', 0)

    pv_full = [min(1.0, p / 1000.0) * pv_cap * pv_eff for p in pv_raw[:total_steps]]
    wind_full = [min(1.0, (w / 3.6) / 12.0) * wind_cap * wind_eff for w in wind_raw[:total_steps]]
    load_ratio_full = [max(0.0, min(1.0, l)) for l in load_raw[:total_steps]]
    load_kw_full = [l * load_base for l in load_ratio_full]
    buy_full = buy_raw[:total_steps]; sell_full = sell_raw[:total_steps]
    uncontrolled_grid_full = [
        max(0.0, load_kw_full[i] - pv_full[i] - wind_full[i])
        for i in range(total_steps)
    ]
    vpp = cfg.get('vpp', {})
    vpp_enabled = vpp.get('enabled', False)
    vpp_charge_price = vpp.get('charge_price', 0.2)
    vpp_response_cap = vpp.get('response_cap_kw')
    vpp_baseline_full = build_workday_average_baseline(
        uncontrolled_grid_full,
        baseline_days=vpp.get('baseline_days', 5),
        steps_per_day=vpp.get('steps_per_day', 96),
    ) if vpp_enabled else []
    vpp_window_full = build_vpp_window_mask(
        ts_all[:total_steps],
        start_hour=vpp.get('start_hour', 11),
        end_hour=vpp.get('end_hour', 13),
    ) if vpp_enabled else []

    print(f"  {run_steps} steps ({run_steps / 96:.0f} days), t0={start_time}")

    # Load forecaster
    forecast_mode = mpc_cfg.get('forecast_mode', 'lightgbm')
    forecaster = None
    lightgbm_min_history_steps = 96
    forecast_history_prefix = np.array([], dtype=np.float64)
    if forecast_mode == 'lightgbm':
        model_path = _resolve_path(mpc_cfg['model_path'])
        if not os.path.exists(model_path):
            print(f"ERROR: model not found at {model_path}")
            sys.exit(1)
        forecaster = LoadForecaster.load(model_path)
        print(f"  Loaded model: {model_path}")
        lightgbm_min_history_steps = _lightgbm_min_history_steps(forecaster)
        print(f"  LightGBM min history: {lightgbm_min_history_steps} steps")
        forecast_history_prefix, history_label = _load_forecast_history_prefix(cfg, forecaster, load_base)
        if len(forecast_history_prefix) > 0:
            print(f"  Forecast history: {len(forecast_history_prefix)} steps from {history_label}")
    elif forecast_mode in ('naive96', 'naive672'):
        from mpc.solvers.forecasting.baselines import seasonal_naive_t96, seasonal_naive_t672
        _naive = seasonal_naive_t96 if forecast_mode == 'naive96' else seasonal_naive_t672
        _naive_lag = 96 if forecast_mode == 'naive96' else 672

    # Config builder
    billing_days = cost.get('billing_days', 30)
    soc_min = battery.get('soc_min', 0.1)
    soc_max = battery.get('soc_max', 0.9)
    bat_kwh = battery['capacity_kwh']
    bat_kw = battery['discharge_max_kw']
    chg_eff = battery.get('charge_eff', 0.95)
    dis_eff = battery.get('discharge_eff', 0.95)
    c_deg = cost.get('c_deg', 0.05)
    demand_rate = cost.get('demand_rate', 0)
    peak_slack_penalty = mpc_cfg.get('peak_slack_penalty', 5000.0)
    battery_smoothing_enabled = bool(mpc_cfg.get('enable_battery_power_smoothing', False))
    if battery_smoothing_enabled:
        battery_ramp_limit_kw = (
            mpc_cfg.get('battery_ramp_limit_kw')
            or battery.get('ramp_limit_kw')
            or battery.get('battery_ramp_limit_kw')
        )
        battery_smooth_penalty = mpc_cfg.get('battery_smooth_penalty', battery.get('smooth_penalty', 0.0))
        battery_ramp_slack_penalty = mpc_cfg.get('battery_ramp_slack_penalty', battery.get('ramp_slack_penalty', 100.0))
    else:
        battery_ramp_limit_kw = None
        battery_smooth_penalty = 0.0
        battery_ramp_slack_penalty = 100.0
    target_peak_kw = mpc_cfg.get('target_peak_kw')
    target_peak_mode = _resolve_target_peak_mode(mpc_cfg)
    target_peak_source = ''

    def make_cfg(
        soc_now,
        peak_sofar,
        vpp_baseline=None,
        vpp_window=None,
        soc_min_profile=None,
        reserve_request=None,
    ):
        reserve_request = reserve_request or {}
        return BenchmarkConfig(
            battery_capacity_kwh=bat_kwh,
            battery_charge_max_kw=battery['charge_max_kw'],
            battery_discharge_max_kw=bat_kw,
            battery_soc_init=soc_now,
            battery_soc_min=soc_min,
            battery_soc_max=soc_max,
            battery_soc_min_profile=soc_min_profile or [],
            battery_charge_eff=chg_eff,
            battery_discharge_eff=dis_eff,
            battery_ramp_limit_kw=battery_ramp_limit_kw,
            battery_initial_power_kw=mpc_cfg.get('_previous_battery_power_kw'),
            battery_smooth_penalty=battery_smooth_penalty,
            battery_ramp_slack_penalty=battery_ramp_slack_penalty,
            grid_import_max_kw=grid.get('import_max_kw', 5000),
            grid_export_max_kw=grid.get('export_max_kw', 0),
            transformer_capacity_kw=grid.get('transformer_capacity_kw', 5000),
            anti_backflow=grid.get('anti_backflow', True),
            c_deg=c_deg,
            r_demand=demand_rate,
            r_capacity=cost.get('capacity_rate', 0),
            use_milp=True,
            terminal_constraint=False,
            optimization_billing_days=billing_days,
            peak_so_far_kw=peak_sofar,
            optimal_peak_kw=target_peak_kw or 0.0,
            peak_slack_penalty=peak_slack_penalty,
            reserve_check_step=reserve_request.get('check_step'),
            reserve_target_soc=reserve_request.get('target_soc', 0.0),
            reserve_slack_penalty=mpc_cfg.get('reserve_slack_penalty', 0.0),
            vpp_enabled=vpp_enabled,
            vpp_charge_price=vpp_charge_price,
            vpp_baseline_kw=vpp_baseline or [],
            vpp_window_mask=vpp_window or [],
            vpp_response_cap_kw=vpp_response_cap,
        )

    # Build forecast covariates for entire simulation period
    fc_cov_cfg = cfg.get('forecast_covariates', {})
    all_covariates = None
    if fc_cov_cfg:
        print("Building forecast covariates...")
        # Timestamps for the entire simulation + lookback
        lookback_steps = max(_naive_lag if forecast_mode in ('naive96', 'naive672') else 96, 96)
        cov_start = ts_all[max(0, start_step - lookback_steps)]
        cov_end_idx = min(start_step + run_steps + horizon, total_steps - 1)
        cov_end = ts_all[cov_end_idx]
        all_ts = pd.date_range(start=cov_start, end=cov_end, freq='15min')
        cov_dfs = []

        # Holiday calendar
        holiday_file = fc_cov_cfg.get('holiday_calendar_file')
        if holiday_file:
            holiday_path = _resolve_path(holiday_file)
            if os.path.exists(holiday_path):
                sheet = fc_cov_cfg.get('holiday_calendar_sheet', 'calendar')
                cal_df = pd.read_excel(holiday_path, sheet_name=sheet)
                cal_cov = build_calendar_covariates(all_ts, public_calendar=cal_df)
                cov_dfs.append(cal_cov)
                print(f"  Loaded holiday calendar: {holiday_path}")

        # Future weather
        weather_file = fc_cov_cfg.get('future_weather_file')
        if weather_file:
            weather_path = _resolve_path(weather_file)
            if os.path.exists(weather_path):
                sheet = fc_cov_cfg.get('future_weather_sheet', 'weather')
                weather_raw = pd.read_excel(weather_path, sheet_name=sheet)

                # Find time column
                time_col = None
                for kw in ['time', 'time', 'date', 'timestamp']:
                    for col in weather_raw.columns:
                        if kw.lower() in str(col).lower():
                            time_col = col
                            break
                    if time_col:
                        break

                if time_col:
                    weather_raw[time_col] = pd.to_datetime(weather_raw[time_col])
                    weather_raw = weather_raw.rename(columns={time_col: 'time'})

                    # Find numeric weather columns
                    numeric_cols = [c for c in weather_raw.columns
                                    if c != 'time' and pd.api.types.is_numeric_dtype(weather_raw[c])]
                    if numeric_cols:
                        weather_cov = align_covariates(all_ts, weather_raw, numeric_cols)
                        cov_dfs.append(weather_cov)
                        print(f"  Loaded weather: {weather_path} (cols: {numeric_cols})")

        if cov_dfs:
            all_covariates = cov_dfs[0]
            for other in cov_dfs[1:]:
                new_cols = [c for c in other.columns if c not in all_covariates.columns]
                if new_cols:
                    all_covariates = all_covariates.merge(
                        other[['time'] + new_cols], on='time', how='left'
                    )
            print(f"  Covariate columns: {[c for c in all_covariates.columns if c != 'time']}")

    # ── MPC main loop ──
    print(f"MPC: {run_steps // 96}d, horizon={horizon}, forecast={forecast_mode}")
    if target_peak_mode == 'manual':
        target_peak_kw = _parse_manual_target_peak(mpc_cfg)
        target_peak_source = 'manual target_peak_kw'
    elif target_peak_mode == 'oracle_full_period':
        from mpc.solvers.mpc.peak_optimizer import find_optimal_peak
        base_peak_cfg = make_cfg(battery.get('soc_init', 0.5), 0.0)
        target_peak_kw = find_optimal_peak(
            pv_full[start_step:start_step + run_steps],
            wind_full[start_step:start_step + run_steps],
            load_kw_full[start_step:start_step + run_steps],
            buy_full[start_step:start_step + run_steps],
            sell_full[start_step:start_step + run_steps],
            base_peak_cfg,
            max(1, run_steps // 96),
        )
        target_peak_source = 'oracle full-period actual data'
    elif target_peak_mode == 'forecast_month_plan':
        from mpc.solvers.mpc.peak_optimizer import find_optimal_peak
        plan_cfg = mpc_cfg.get('target_peak_forecast') or {}
        base_peak_cfg = make_cfg(battery.get('soc_init', 0.5), 0.0)
        fpv, fwind, fload, fbuy, fsell = _load_target_peak_forecast_inputs(
            plan_cfg,
            sheets_cfg,
            load_base,
            device,
            (
                buy_full[start_step:start_step + run_steps],
                sell_full[start_step:start_step + run_steps],
            ),
        )
        target_peak_kw = find_optimal_peak(
            fpv,
            fwind,
            fload,
            fbuy,
            fsell,
            base_peak_cfg,
            max(1, run_steps // 96),
        )
        data_file = plan_cfg.get('data_file') or plan_cfg.get('file')
        target_peak_source = f"forecast month plan: {data_file}"
    print(
        f"  target_peak={target_peak_kw:.1f}kW "
        f"mode={target_peak_mode} source={target_peak_source} "
        f"slack_penalty={peak_slack_penalty:.0f}"
    )
    if battery_smoothing_enabled:
        print(
            f"  battery smoothing: ramp_limit={battery_ramp_limit_kw}kW "
            f"smooth_penalty={battery_smooth_penalty} "
            f"ramp_slack_penalty={battery_ramp_slack_penalty}"
        )
    soc = battery.get('soc_init', 0.5)
    running_cost = 0.0; peak = 0.0; peak_sofar = 0.0
    traj = []; fc_errors = []
    under_error_history = []
    previous_battery_power_kw = float(mpc_cfg.get('initial_battery_power_kw', 0.0))
    progress_interval_steps = max(1, int(mpc_cfg.get('progress_interval_steps', 24)))
    t_start = time.time()

    for t in range(run_steps):
        abs_t = start_step + t
        end = min(t + horizon, run_steps)
        abs_end = start_step + end

        # Forecast
        if forecast_mode == 'file':
            load_fc = load_kw_full[abs_t:abs_end]
        elif forecast_mode == 'lightgbm' and forecaster:
            hist = _compose_forecast_history(load_ratio_full, abs_t, forecast_history_prefix)
            if len(hist) < lightgbm_min_history_steps:
                raise RuntimeError(
                    f"LightGBM history is too short at step {t}: "
                    f"{len(hist)} < {lightgbm_min_history_steps}. "
                    "Configure forecast_history with pre-run historical load data; "
                    "MPC no longer falls back to future actual load."
                )
            st = start_time + timedelta(minutes=15 * t)

            # Extract future covariates for this prediction horizon
            future_cov = None
            if all_covariates is not None:
                future_ts = pd.date_range(start=st, periods=end - t, freq='15min')
                mask = all_covariates['time'].isin(future_ts)
                if mask.any():
                    future_cov = all_covariates[mask].reset_index(drop=True)

            fc_r = forecaster.predict(hist, st, end - t,
                                       future_covariates=future_cov)
            load_fc = [max(0, r * load_base) for r in fc_r]
        elif forecast_mode in ('naive96', 'naive672') and t >= _naive_lag:
            hist = np.array(load_kw_full[:abs_t], dtype=np.float64)
            load_fc = list(_naive(hist, end - t))
        else:
            load_fc = load_kw_full[abs_t:abs_end]

        # Solve MILP
        reserve_energy_enabled = mpc_cfg.get('reserve_energy_enabled', False)
        reserve_enabled = mpc_cfg.get('reserve_soc_enabled', False) and not reserve_energy_enabled
        effective_target_kw = max(target_peak_kw or 0.0, peak_sofar)
        reserve_soc_profile = build_dynamic_reserve_soc(
            load_fc=load_fc,
            pv_kw=pv_full[abs_t:abs_end],
            wind_kw=wind_full[abs_t:abs_end],
            effective_target_kw=effective_target_kw,
            base_soc=mpc_cfg.get('reserve_soc_normal', soc_min),
            risk_soc=mpc_cfg.get('reserve_soc_risk', 0.25),
            high_soc=mpc_cfg.get('reserve_soc_high', 0.35),
            risk_margin_kw=mpc_cfg.get('reserve_margin_risk_kw', 120.0),
            high_margin_kw=mpc_cfg.get('reserve_margin_high_kw', 60.0),
            enabled=reserve_enabled,
            current_soc=soc,
            charge_max_kw=battery['charge_max_kw'],
            capacity_kwh=bat_kwh,
            charge_eff=chg_eff,
            dt_hours=DT,
            soc_max=soc_max,
        )
        lookback_steps = int(mpc_cfg.get('reserve_error_lookback_days', 3) * 96)
        recent_under_errors = under_error_history[-lookback_steps:] if lookback_steps > 0 else under_error_history
        risk_margin_kw = estimate_under_forecast_margin(
            recent_under_errors,
            quantile=mpc_cfg.get('reserve_error_quantile', 0.90),
            warmup_kw=mpc_cfg.get('reserve_error_warmup_kw', 100.0),
            min_kw=mpc_cfg.get('reserve_error_min_kw', 30.0),
            max_kw=mpc_cfg.get('reserve_error_max_kw', 150.0),
            min_samples=mpc_cfg.get('reserve_error_min_samples', 96),
        )
        reserve_request = build_reserve_energy_request(
            load_fc=load_fc,
            pv_kw=pv_full[abs_t:abs_end],
            wind_kw=wind_full[abs_t:abs_end],
            effective_target_kw=effective_target_kw,
            risk_margin_kw=risk_margin_kw,
            soc_min=soc_min,
            soc_max=soc_max,
            capacity_kwh=bat_kwh,
            discharge_eff=dis_eff,
            dt_hours=DT,
            reserve_lead_hours=mpc_cfg.get('reserve_lead_hours', 2.0),
            enabled=reserve_energy_enabled,
        )
        vpp_base_h = vpp_baseline_full[abs_t:abs_end] if vpp_enabled else None
        vpp_window_h = vpp_window_full[abs_t:abs_end] if vpp_enabled else None
        mpc_cfg['_previous_battery_power_kw'] = previous_battery_power_kw
        r = solve_benchmark(pv_full[abs_t:abs_end], wind_full[abs_t:abs_end], load_fc,
                            buy_full[abs_t:abs_end], sell_full[abs_t:abs_end],
                            config=make_cfg(
                                soc,
                                peak_sofar,
                                vpp_base_h,
                                vpp_window_h,
                                reserve_soc_profile,
                                reserve_request,
                            ),
                            verbose=False)
        bp = r.battery_power[0] if r.battery_power else 0.0
        vpp_window_now = vpp_window_full[abs_t] if vpp_enabled else False

        # Deviation compensation
        dev = load_kw_full[abs_t] - load_fc[0]
        under_error_history.append(max(0.0, dev))
        bp = apply_realtime_load_compensation(
            battery_power_kw=bp,
            load_error_kw=dev,
            discharge_max_kw=bat_kw,
            charge_max_kw=battery['charge_max_kw'],
            previous_battery_power_kw=previous_battery_power_kw,
            ramp_limit_kw=battery_ramp_limit_kw,
            current_soc=soc,
            soc_min=soc_min,
            soc_max=soc_max,
            capacity_kwh=bat_kwh,
            charge_eff=chg_eff,
            discharge_eff=dis_eff,
            dt_hours=DT,
            vpp_window_now=vpp_window_now,
            vpp_discharge_compensation=mpc_cfg.get('vpp_realtime_discharge_compensation', True),
        )

        chg, dis = max(0, -bp), max(0, bp)
        gi = max(0, load_kw_full[abs_t] - pv_full[abs_t] - wind_full[abs_t] - bp)
        go = max(0, -(load_kw_full[abs_t] - pv_full[abs_t] - wind_full[abs_t] - bp))
        vpp_resp, vpp_benefit = compute_vpp_step_benefit(
            gi,
            vpp_baseline_full[abs_t] if vpp_enabled else 0.0,
            buy_full[abs_t],
            vpp_charge_price,
            vpp_window_now,
            vpp_response_cap,
            DT,
        )
        step_cost = (
            buy_full[abs_t] * gi * DT
            - sell_full[abs_t] * go * DT
            + c_deg * (chg + dis) * DT
            - vpp_benefit
        )
        # arbitrage: discharging avoids buying at buy_price, charging pays buy_price
        arbitrage = bp * buy_full[abs_t] * DT
        traj.append([pv_full[abs_t], wind_full[abs_t], load_kw_full[abs_t],
                     soc, bp, gi, go, buy_full[abs_t], sell_full[abs_t], step_cost,
                     vpp_resp, vpp_benefit, arbitrage])
        fc_errors.append(abs(dev))
        soc += (chg_eff * chg - dis / dis_eff) * DT / bat_kwh
        soc = max(soc_min, min(soc_max, soc))
        previous_battery_power_kw = bp
        running_cost += step_cost; peak_sofar = max(peak_sofar, gi); peak = max(peak, gi)
        if t == 0 or (t + 1) % progress_interval_steps == 0 or (t + 1) == run_steps:
            elapsed_now = time.time() - t_start
            print(
                f"  step {t + 1}/{run_steps} "
                f"day={(t + 1) / 96:.2f}/{run_steps / 96:.0f}: "
                f"SOC={soc:.3f} peak={peak_sofar:.1f}kW cost={running_cost:.0f} "
                f"bat={bp:.1f} grid={gi:.1f} load={float(load_kw_full[abs_t]):.1f} "
                f"pv={float(pv_full[abs_t]):.1f} arb={arbitrage:.3f} "
                f"elapsed={elapsed_now:.0f}s",
                flush=True,
            )

    total_demand = demand_rate * peak * run_steps * DT / 24 / 30
    ctrl = running_cost + total_demand
    fc_mae = float(np.mean(fc_errors))
    elapsed_mpc = time.time() - t_start
    print(f"  MPC: peak={peak:.1f}kW demand={total_demand:.0f} basic={running_cost:.0f} "
          f"ctrl={ctrl:.0f} FC_MAE={fc_mae:.2f}kW [{elapsed_mpc:.0f}s]")

    # ── MILP benchmark ──
    print("Running MILP benchmark...")
    pv_kw = pv_full[start_step:start_step + run_steps]
    wind_kw = wind_full[start_step:start_step + run_steps]
    load_kw = load_kw_full[start_step:start_step + run_steps]
    buy = buy_full[start_step:start_step + run_steps]
    sell = sell_full[start_step:start_step + run_steps]
    vpp_base = vpp_baseline_full[start_step:start_step + run_steps] if vpp_enabled else []
    vpp_window = vpp_window_full[start_step:start_step + run_steps] if vpp_enabled else []

    bench_cfg = BenchmarkConfig(
        battery_capacity_kwh=bat_kwh,
        battery_charge_max_kw=battery['charge_max_kw'],
        battery_discharge_max_kw=bat_kw,
        battery_soc_init=battery.get('soc_init', 0.5),
        battery_soc_min=soc_min,
        battery_soc_max=soc_max,
        battery_charge_eff=chg_eff,
        battery_discharge_eff=dis_eff,
        grid_import_max_kw=grid.get('import_max_kw', 5000),
        grid_export_max_kw=grid.get('export_max_kw', 0),
        transformer_capacity_kw=grid.get('transformer_capacity_kw', 5000),
        anti_backflow=grid.get('anti_backflow', True),
        c_deg=c_deg,
        r_demand=demand_rate,
        r_capacity=cost.get('capacity_rate', 0),
        use_milp=True,
        terminal_constraint=False,
        optimization_billing_days=billing_days,
        vpp_enabled=vpp_enabled,
        vpp_charge_price=vpp_charge_price,
        vpp_baseline_kw=vpp_base,
        vpp_window_mask=vpp_window,
        vpp_response_cap_kw=vpp_response_cap,
    )
    bench_r = solve_benchmark(pv_kw, wind_kw, load_kw, buy, sell, config=bench_cfg, verbose=False)
    bench_demand = demand_rate * bench_r.peak_demand_kw * run_steps * DT / 24 / 30
    bench_basic = (
        bench_r.purchase_cost
        - bench_r.export_revenue
        - bench_r.vpp_benefit
        + bench_r.degradation_cost
    )
    bench_ctrl = bench_basic + bench_demand
    achieve = (bench_ctrl / ctrl * 100) if ctrl > 0 else 0.0
    print(f"  MILP: peak={bench_r.peak_demand_kw:.1f}kW basic={bench_basic:.0f} "
          f"demand={bench_demand:.0f} ctrl={bench_ctrl:.0f}")
    print(f"  Achievement rate = {achieve:.1f}%")

    # ── Write Excel ──
    print("Writing Excel...")
    hf = Font(bold=True)
    hfl = PatternFill("solid", fgColor="DDEEFF")
    ha = Alignment(horizontal="center")
    wb = Workbook()

    def set_h(ws, hdrs):
        for c, h in enumerate(hdrs, 1):
            cell = ws.cell(row=1, column=c, value=h)
            cell.font = hf; cell.fill = hfl; cell.alignment = ha

    # Sheet 1: 15min_trajectory
    ws1 = wb.active; ws1.title = '15min_trajectory'
    set_h(ws1, ['时间', '光伏出力(kW)', '风电出力(kW)', '负荷功率(kW)', 'SOC',
                '电池功率(kW)', '电网功率(kW)', '购电价(元/kWh)', '售电价(元/kWh)',
                '互济结算功率(kW)', '互济收益(元)', '套利收益(元)'])
    for i, d in enumerate(traj):
        ts = start_time + timedelta(minutes=15 * i)
        ws1.append([ts.strftime('%Y-%m-%d %horizon:%M:%S'), round(d[0], 4), round(d[1], 4),
                    round(d[2], 4), round(d[3], 4), round(d[4], 4), round(d[5] - d[6], 4),
                    round(d[7], 6), round(d[8], 6), round(d[10], 4), round(d[11], 4),
                    round(d[12], 4)])
    for c in range(1, 13):
        ws1.column_dimensions[get_column_letter(c)].width = 20

    # Sheet 2: daily_summary
    ws2 = wb.create_sheet('daily_summary')
    set_h(ws2, ['天数', '日期', '起始SOC', '结束SOC', '最低SOC', '最高SOC',
                '光伏(kWh)', '风电(kWh)', '负荷(kWh)', '购电费(元)', '售电收益(元)',
                '互济收益(元)', '套利收益(元)', '衰减(元)', '净成本(元)'])
    for day in range(run_steps // 96):
        s, e = day * 96, (day + 1) * 96
        dd = traj[s:e]
        day_soc = [r[3] for r in dd]
        day_bat = [r[4] for r in dd]
        day_gi = [r[5] for r in dd]
        day_go = [r[6] for r in dd]
        day_vpp = [r[11] for r in dd]
        day_arb = [r[12] for r in dd]
        dpv = sum(r[0] for r in dd) * DT
        dwind = sum(r[1] for r in dd) * DT
        dload = sum(r[2] for r in dd) * DT
        dbuy = sum(dd[i][7] * day_gi[i] * DT for i in range(96))
        brev = sum(dd[i][8] * day_go[i] * DT for i in range(96))
        dvpp = sum(day_vpp)
        darb = sum(day_arb)
        ddeg = sum(c_deg * (max(0, -b) + max(0, b)) * DT for b in day_bat)
        ws2.append([day + 1, (start_time + timedelta(days=day)).strftime('%Y-%m-%d'),
                    round(day_soc[0], 4), round(day_soc[-1], 4),
                    round(min(day_soc), 4), round(max(day_soc), 4),
                    round(dpv, 1), round(dwind, 1), round(dload, 1),
                    round(dbuy, 2), round(brev, 2), round(dvpp, 2),
                    round(darb, 2), round(ddeg, 2),
                    round(dbuy - brev - dvpp + ddeg, 2)])
    for c in range(1, 16):
        ws2.column_dimensions[get_column_letter(c)].width = 16

    # Sheet 3: cost_summary
    ws3 = wb.create_sheet('cost_summary')
    set_h(ws3, ['指标', 'MPC 数值', 'MILP 基准', '单位/说明'])
    tp = sum(d[5] * d[7] * DT for d in traj)
    trv = sum(d[6] * d[8] * DT for d in traj)
    tvpp = sum(d[11] for d in traj)
    tarb = sum(d[12] for d in traj)
    td_deg = sum(c_deg * (max(0, -d[4]) + max(0, d[4])) * DT for d in traj)
    rows = [
        ('购电费(元)', round(tp, 2), round(bench_r.purchase_cost, 2), ''),
        ('售电收益(元)', round(trv, 2), round(bench_r.export_revenue, 2), ''),
        ('互济收益(元)', round(tvpp, 2), round(bench_r.vpp_benefit, 2), ''),
        ('套利收益(元)', round(tarb, 2), '', '放电省购电 - 充电花购电'),
        ('储能衰减(元)', round(td_deg, 2), round(bench_r.degradation_cost, 2), ''),
        ('需量费(元)', round(total_demand, 2), round(bench_demand, 2), ''),
        ('可控成本(元)', round(ctrl, 2), round(bench_ctrl, 2), ''),
        ('峰值需量(kW)', round(peak, 1), round(bench_r.peak_demand_kw, 1), ''),
        ('目标峰值(kW)', round(target_peak_kw, 1), '', ''),
        ('目标峰值模式', target_peak_mode, '', target_peak_source),
        ('预测MAE(kW)', round(fc_mae, 2), '', ''),
        ('成本达成率', f'{achieve:.1f}%', '', ''),
    ]
    for i, row in enumerate(rows):
        for c, v in enumerate(row, 1):
            ws3.cell(row=i + 2, column=c, value=v)
    ws3.column_dimensions['A'].width = 22
    for c in ['B', 'C', 'D']:
        ws3.column_dimensions[c].width = 16

    os.makedirs(os.path.dirname(out_path) or '.', exist_ok=True)
    wb.save(out_path)
    print(f"Done: {out_path}")


if __name__ == '__main__':
    main()
