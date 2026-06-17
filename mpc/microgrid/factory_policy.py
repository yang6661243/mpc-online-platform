"""Factory policy cost report CLI.

This is a legacy cost-report entrypoint. The production optimization flow uses
``microgrid.forecaster``, ``microgrid.solver``, and ``microgrid.mpc`` with the
profiled configs under ``examples/``.
"""

from __future__ import annotations

import argparse
import os
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import pandas as pd
import yaml
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[2]
DT_HOURS = 0.25


@dataclass
class FactoryPolicyReport:
    trajectory: pd.DataFrame
    daily_summary: pd.DataFrame
    cost_summary: dict[str, float]
    alignment: dict[str, Any]


def _resolve_path(path: str | os.PathLike[str]) -> Path:
    p = Path(path)
    if p.is_absolute():
        return p
    return ROOT / p


def _select_series(df: pd.DataFrame, spec: Any, label: str) -> pd.Series:
    if isinstance(spec, int):
        idx = spec - 1
        if idx < 0 or idx >= len(df.columns):
            raise ValueError(f"{label} column index {spec} is outside available columns")
        return df.iloc[:, idx]

    if isinstance(spec, str):
        if spec in df.columns:
            return df[spec]
        spec_l = spec.lower()
        matches = [c for c in df.columns if spec_l in str(c).lower()]
        if matches:
            return df[matches[0]]

    if isinstance(spec, list):
        for item in spec:
            try:
                return _select_series(df, item, label)
            except ValueError:
                continue

    raise ValueError(f"Cannot find {label} column from spec {spec!r}; columns={list(df.columns)!r}")


def _read_sheet(path: Path, sheet: str | None = None) -> pd.DataFrame:
    if sheet:
        return pd.read_excel(path, sheet_name=sheet)
    return pd.read_excel(path)


def _read_scenario(cfg: dict[str, Any]) -> pd.DataFrame:
    path = _resolve_path(cfg["file"])
    sheets = cfg.get("sheets", {})
    load_base = float(cfg.get("load_base_kw", cfg.get("load_base", 1000.0)))
    pv_capacity = float(cfg.get("pv_capacity_kw", 1000.0))
    pv_efficiency = float(cfg.get("pv_efficiency", 1.0))

    load_df = _read_sheet(path, sheets.get("load", "load_apr"))
    pv_df = _read_sheet(path, sheets.get("pv_wind", "pv_apr"))
    price_df = _read_sheet(path, sheets.get("price", "price"))

    load = pd.DataFrame(
        {
            "time": pd.to_datetime(load_df.iloc[:, 0], errors="coerce"),
            "load_kw": pd.to_numeric(_select_series(load_df, cfg.get("load_col", ["load", "demand", 2]), "load"), errors="coerce")
            * load_base,
        }
    )
    pv_raw = pd.to_numeric(
        _select_series(pv_df, cfg.get("pv_col", ["irradiance", "solar", 2]), "pv"),
        errors="coerce",
    )
    pv = pd.DataFrame(
        {
            "time": pd.to_datetime(pv_df.iloc[:, 0], errors="coerce"),
            "pv_kw": pv_raw.clip(lower=0.0, upper=1000.0) / 1000.0 * pv_capacity * pv_efficiency,
        }
    )
    price = pd.DataFrame(
        {
            "time": pd.to_datetime(price_df.iloc[:, 0], errors="coerce"),
            "buy_price": pd.to_numeric(
                _select_series(price_df, cfg.get("buy_price_col", ["buy_price", "购电价", 2]), "buy_price"),
                errors="coerce",
            ),
            "sell_price": pd.to_numeric(
                _select_series(price_df, cfg.get("sell_price_col", ["sell_price", "售电价", 3]), "sell_price"),
                errors="coerce",
            ),
        }
    )

    scenario = load.merge(pv, on="time", how="inner").merge(price, on="time", how="inner")
    return scenario.dropna(subset=["time"]).sort_values("time")


def _read_actual(cfg: dict[str, Any]) -> pd.DataFrame:
    path = _resolve_path(cfg["file"])
    df = _read_sheet(path, cfg.get("sheet"))
    cols = cfg.get("columns", {})
    time_spec = cols.get("time", ["time", "时间", 1])
    grid_spec = cols.get("grid_power_kw", ["grid", "电网", "变压器"])
    battery_spec = cols.get("battery_power_kw")
    soc_spec = cols.get("soc")

    actual = pd.DataFrame(
        {
            "time": pd.to_datetime(_select_series(df, time_spec, "actual time"), errors="coerce"),
            "grid_power_kw": pd.to_numeric(_select_series(df, grid_spec, "actual grid_power_kw"), errors="coerce"),
        }
    )

    if battery_spec is not None:
        actual["battery_power_kw"] = pd.to_numeric(
            _select_series(df, battery_spec, "actual battery_power_kw"),
            errors="coerce",
        ).fillna(0.0)
    else:
        actual["battery_power_kw"] = 0.0

    if soc_spec is not None:
        actual["soc"] = pd.to_numeric(_select_series(df, soc_spec, "actual soc"), errors="coerce")
    else:
        actual["soc"] = pd.NA

    return actual.dropna(subset=["time", "grid_power_kw"]).sort_values("time")


def _apply_window(df: pd.DataFrame, cfg: dict[str, Any]) -> pd.DataFrame:
    start = cfg.get("start")
    end = cfg.get("end")
    out = df
    if start:
        out = out[out["time"] >= pd.Timestamp(start)]
    if end:
        out = out[out["time"] <= pd.Timestamp(end)]
    return out.copy()


def _add_battery_energy_columns(df: pd.DataFrame, sign: str) -> pd.DataFrame:
    out = df.copy()
    battery = out["battery_power_kw"].astype(float)
    if sign == "positive_discharge":
        out["charge_kw"] = (-battery).clip(lower=0.0)
        out["discharge_kw"] = battery.clip(lower=0.0)
    else:
        out["charge_kw"] = battery.clip(lower=0.0)
        out["discharge_kw"] = (-battery).clip(lower=0.0)
    return out


def _validate_alignment(df: pd.DataFrame, evaluation_cfg: dict[str, Any]) -> dict[str, Any]:
    freq_minutes = int(evaluation_cfg.get("freq_minutes", 15))
    if df.empty:
        raise ValueError("No aligned data points found for the requested evaluation window")

    expected = pd.date_range(df["time"].min(), df["time"].max(), freq=f"{freq_minutes}min")
    actual_times = pd.DatetimeIndex(df["time"])
    missing = expected.difference(actual_times)
    duplicates = int(actual_times.duplicated().sum())
    if len(missing) or duplicates:
        raise ValueError(
            f"Aligned data is not a continuous {freq_minutes} minute series: "
            f"missing={len(missing)}, duplicates={duplicates}"
        )

    return {
        "start": df["time"].min(),
        "end": df["time"].max(),
        "aligned_points": int(len(df)),
        "missing_points": int(len(missing)),
        "freq_minutes": freq_minutes,
    }


def build_factory_policy_report(cfg: dict[str, Any]) -> FactoryPolicyReport:
    scenario = _apply_window(_read_scenario(cfg["scenario"]), cfg.get("evaluation", {}))
    actual = _apply_window(_read_actual(cfg["actual"]), cfg.get("evaluation", {}))
    merged = scenario.merge(actual, on="time", how="inner").sort_values("time").reset_index(drop=True)
    alignment = _validate_alignment(merged, cfg.get("evaluation", {}))

    sign = cfg.get("actual", {}).get("battery_power_sign", "positive_charge")
    traj = _add_battery_energy_columns(merged, sign)
    cost_cfg = cfg.get("cost", {})
    c_deg = float(cost_cfg.get("c_deg", 0.05))
    demand_rate = float(cost_cfg.get("demand_rate", 0.0))
    capacity_rate = float(cost_cfg.get("capacity_rate", 0.0))
    transformer_capacity = float(cost_cfg.get("transformer_capacity_kw", 0.0))
    billing_days = float(cost_cfg.get("billing_days", 30.0))
    eval_days = len(traj) * DT_HOURS / 24.0

    traj["grid_import_kw"] = traj["grid_power_kw"].clip(lower=0.0)
    traj["grid_export_kw"] = (-traj["grid_power_kw"]).clip(lower=0.0)
    traj["purchase_cost"] = traj["grid_import_kw"] * traj["buy_price"] * DT_HOURS
    traj["export_revenue"] = traj["grid_export_kw"] * traj["sell_price"] * DT_HOURS
    traj["degradation_cost"] = (traj["charge_kw"] + traj["discharge_kw"]) * c_deg * DT_HOURS
    traj["net_cost"] = traj["purchase_cost"] - traj["export_revenue"] + traj["degradation_cost"]

    daily_rows = []
    for idx, (date, day) in enumerate(traj.groupby(traj["time"].dt.date), 1):
        daily_rows.append(
            {
                "day": idx,
                "date": str(date),
                "start_soc": _first_numeric(day["soc"]),
                "end_soc": _last_numeric(day["soc"]),
                "min_soc": _min_numeric(day["soc"]),
                "max_soc": _max_numeric(day["soc"]),
                "load_energy_kwh": float(day["load_kw"].sum() * DT_HOURS),
                "charge_energy_kwh": float(day["charge_kw"].sum() * DT_HOURS),
                "discharge_energy_kwh": float(day["discharge_kw"].sum() * DT_HOURS),
                "purchase_cost": float(day["purchase_cost"].sum()),
                "degradation_cost": float(day["degradation_cost"].sum()),
                "net_cost": float(day["net_cost"].sum()),
            }
        )
    daily = pd.DataFrame(daily_rows)

    purchase_cost = float(traj["purchase_cost"].sum())
    export_revenue = float(traj["export_revenue"].sum())
    degradation_cost = float(traj["degradation_cost"].sum())
    peak_demand_kw = float(traj["grid_import_kw"].max())
    demand_charge = demand_rate * peak_demand_kw * eval_days / billing_days
    capacity_charge = capacity_rate * transformer_capacity * eval_days / billing_days
    comprehensive_cost = purchase_cost - export_revenue + degradation_cost + demand_charge + capacity_charge

    costs = {
        "purchase_cost": round(purchase_cost, 6),
        "export_revenue": round(export_revenue, 6),
        "degradation_cost": round(degradation_cost, 6),
        "demand_charge": round(demand_charge, 6),
        "capacity_charge": round(capacity_charge, 6),
        "peak_demand_kw": round(peak_demand_kw, 6),
        "comprehensive_cost": round(comprehensive_cost, 6),
    }
    return FactoryPolicyReport(traj, daily, costs, alignment)


def _first_numeric(series: pd.Series) -> float | None:
    valid = series.dropna()
    return None if valid.empty else float(valid.iloc[0])


def _last_numeric(series: pd.Series) -> float | None:
    valid = series.dropna()
    return None if valid.empty else float(valid.iloc[-1])


def _min_numeric(series: pd.Series) -> float | None:
    valid = series.dropna()
    return None if valid.empty else float(valid.min())


def _max_numeric(series: pd.Series) -> float | None:
    valid = series.dropna()
    return None if valid.empty else float(valid.max())


def write_factory_policy_excel(report: FactoryPolicyReport, output_path: str | os.PathLike[str]) -> None:
    output = _resolve_path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)

    trajectory = report.trajectory[
        [
            "time",
            "pv_kw",
            "load_kw",
            "soc",
            "battery_power_kw",
            "grid_power_kw",
            "buy_price",
            "sell_price",
        ]
    ].copy()
    trajectory.columns = [
        "时间",
        "光伏功率(kW)_实测",
        "负荷功率(kW)_实测",
        "SOC_实测",
        "电池功率(kW)_实测",
        "电网功率(kW)_实测",
        "购电价(元/kWh)",
        "售电价(元/kWh)",
    ]

    daily = report.daily_summary.copy()
    daily.columns = [
        "天数",
        "日期",
        "开始SOC",
        "结束SOC",
        "最低SOC",
        "最高SOC",
        "用电量(kWh)",
        "充电量(kWh)",
        "放电量(kWh)",
        "购电费(元)",
        "衰减(元)",
        "净成本(元)",
    ]

    cost_rows = [
        ("购电费", report.cost_summary["purchase_cost"]),
        ("售电收益", report.cost_summary["export_revenue"]),
        ("电池衰减", report.cost_summary["degradation_cost"]),
        ("需量费", report.cost_summary["demand_charge"]),
        ("容量电费", report.cost_summary["capacity_charge"]),
        ("峰值需量(kW)", report.cost_summary["peak_demand_kw"]),
        ("综合成本", report.cost_summary["comprehensive_cost"]),
        ("目标函数J", report.cost_summary["comprehensive_cost"]),
    ]
    cost = pd.DataFrame(cost_rows, columns=["指标", "数值(元)"])

    alignment_rows = [(k, v) for k, v in report.alignment.items()]
    alignment = pd.DataFrame(alignment_rows, columns=["字段", "值"])

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        trajectory.to_excel(writer, sheet_name="15min_trajectory", index=False)
        daily.to_excel(writer, sheet_name="daily_summary", index=False)
        cost.to_excel(writer, sheet_name="cost_summary", index=False)
        alignment.to_excel(writer, sheet_name="alignment_summary", index=False)

    _format_workbook(output)


def _format_workbook(path: Path) -> None:
    wb = load_workbook(path)
    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="DDEEFF")
    center = Alignment(horizontal="center")
    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
        for col_idx in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = 22
        if ws.max_column and ws.max_row > 1:
            for cell in ws["A"][1:]:
                if ws.title in {"15min_trajectory", "alignment_summary"}:
                    cell.number_format = "yyyy-mm-dd hh:mm:ss"
    wb.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Calculate factory policy cost report")
    parser.add_argument("--config", required=True, help="Path to factory policy YAML config")
    args = parser.parse_args()

    config_path = _resolve_path(args.config)
    if not config_path.exists():
        print(f"ERROR: config file not found: {config_path}")
        sys.exit(1)

    with open(config_path, encoding="utf-8") as f:
        cfg = yaml.safe_load(f)

    report = build_factory_policy_report(cfg)
    output = cfg.get("output", "outputs/factory_policy.xlsx")
    write_factory_policy_excel(report, output)
    print(f"Aligned points: {report.alignment['aligned_points']}")
    print(f"Evaluation: {report.alignment['start']} to {report.alignment['end']}")
    print(f"Comprehensive cost: {report.cost_summary['comprehensive_cost']:.2f}")
    print(f"Done: {_resolve_path(output)}")


if __name__ == "__main__":
    main()
