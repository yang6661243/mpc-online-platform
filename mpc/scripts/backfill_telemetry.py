#!/usr/bin/env python3
"""
将 Excel 15min_trajectory 数据回填到 telemetry_15min 表（工厂策略数据）。

工厂电网功率 = 负荷功率 - 光伏出力（无电池调度的原始状态）
工厂电池功率 = 0（工厂不调度电池）
工厂 SOC = 固定初始值（不调度则不变化）

用法：
  python scripts/data/backfill_telemetry.py outputs/xxx.xlsx --plant-id hehong_huajin
  python scripts/data/backfill_telemetry.py outputs/xxx.xlsx --plant-id hehong_huajin --dry-run
  python scripts/data/backfill_telemetry.py outputs/xxx.xlsx --plant-id hehong_huajin --soc-start 0.5
"""

from __future__ import annotations

import argparse
import sys
from datetime import datetime, timedelta
from pathlib import Path

import openpyxl

PROJECT_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(PROJECT_ROOT))

from api.database import create_session_factory
from api.models import Telemetry15Min
from sqlalchemy import select

DATE_FORMATS = [
    "%Y-%m-%d %H:%M:%S",
    "%Y/%m/%d %H:%M:%S",
    "%Y-%m-%dT%H:%M:%S",
]


def parse_time(value: object) -> datetime:
    if isinstance(value, datetime):
        return value
    if isinstance(value, str):
        for fmt in DATE_FORMATS:
            try:
                return datetime.strptime(value.strip(), fmt)
            except ValueError:
                continue
    raise ValueError(f"无法解析时间: {value!r}")


def read_trajectory(filepath: str) -> list[dict]:
    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    if "15min_trajectory" not in wb.sheetnames:
        wb.close()
        raise ValueError(f"未找到 15min_trajectory sheet，可用: {wb.sheetnames}")

    ws = wb["15min_trajectory"]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    col_idx = {h: i for i, h in enumerate(headers)}

    required = ["时间", "负荷功率(kW)", "光伏出力(kW)"]
    missing = [c for c in required if c not in col_idx]
    if missing:
        raise ValueError(f"Excel 表头缺少必要列: {missing}")

    data = []
    for row in rows[1:]:
        if all(v is None for v in row):
            continue
        try:
            t = parse_time(row[col_idx["时间"]])
        except (ValueError, IndexError):
            continue

        def get_float(col: str) -> float | None:
            val = row[col_idx[col]] if col_idx[col] < len(row) else None
            if val is None:
                return None
            try:
                return float(val)
            except (ValueError, TypeError):
                return None

        load_kw = get_float("负荷功率(kW)")
        pv_kw = get_float("光伏出力(kW)")

        # 工厂电网功率 = 负荷 - 光伏（无电池调度）
        factory_grid = round(load_kw - pv_kw, 6) if load_kw is not None and pv_kw is not None else None

        data.append({
            "time": t,
            "grid_power_kw": factory_grid,
            "load_minus_pv_kw": factory_grid,
        })
    return data


def fmt_month(ts: datetime) -> str:
    return f"{ts.year}年{ts.month}月"


def main():
    parser = argparse.ArgumentParser(description="回填 Excel 工厂策略数据到 telemetry_15min")
    parser.add_argument("excel", help="Excel 文件路径")
    parser.add_argument("--plant-id", default="hehong_huajin", help="电站 ID")
    parser.add_argument("--soc-start", type=float, default=0.5, help="工厂 SOC 初始值（固定不变，默认 0.5）")
    parser.add_argument("--dry-run", action="store_true", help="仅预览，不写入数据库")
    args = parser.parse_args()

    rows = read_trajectory(args.excel)
    if not rows:
        print("没有有效数据行")
        return

    month_str = fmt_month(rows[0]["time"])
    print(f"解析到 {len(rows)} 行 ({month_str})")
    print(f"时间范围: {rows[0]['time']} ~ {rows[-1]['time']}")
    print(f"工厂 SOC 固定值: {args.soc_start}")
    print(f"公式: 工厂电网功率 = 负荷功率 - 光伏出力")

    # 预览几条
    print("\n预览（前 3 行）:")
    for r in rows[:3]:
        print(f"  {r['time']}  工厂电网={r['grid_power_kw']} kW")

    if args.dry_run:
        print(f"\n[Dry-Run] 共 {len(rows)} 行待写入，未实际入库")
        return

    factory = create_session_factory()
    session = factory()

    inserted = 0
    skipped = 0
    for r in rows:
        t = r["time"]
        start_t = t
        end_t = t + timedelta(minutes=15)

        # 去重
        existing = session.scalar(
            select(Telemetry15Min).where(
                Telemetry15Min.plant_id == args.plant_id,
                Telemetry15Min.start_time == start_t,
                Telemetry15Min.end_time == end_t,
            )
        )
        if existing is not None:
            skipped += 1
            continue

        row = Telemetry15Min(
            plant_id=args.plant_id,
            start_time=start_t,
            end_time=end_t,
            grid_power_kw_avg=r["grid_power_kw"],
            grid_power_kw_max=r["grid_power_kw"],
            battery_power_kw_avg=0,       # 工厂不调度电池
            load_minus_pv_kw_avg=r["load_minus_pv_kw"],
            soc_start=args.soc_start,
            soc_end=args.soc_start,
            grid_sample_count=1,
            battery_sample_count=1,
            quality_flag="ok",
        )
        session.add(row)
        inserted += 1

    session.commit()
    session.close()
    print(f"\n完成: 插入 {inserted} 行, 跳过 {skipped} 行（已存在）")
    print(f"工厂策略数据 ({month_str}) 已入库，仪表盘可用")


if __name__ == "__main__":
    main()
