"""Offline gym environment: reads local scenario files, runs battery simulation."""
from __future__ import annotations
import csv
import gymnasium as gym
import numpy as np
from gymnasium import spaces
from openpyxl import load_workbook


def _read_xlsx_col(fname: str, keywords: list[str]) -> list[float]:
    """Read first numeric column from xlsx whose header contains any keyword."""
    wb = load_workbook(fname, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    headers = [str(c.value) if c.value else '' for c in next(ws.iter_rows(min_row=1, max_row=1))]
    target_col = None
    for kw in keywords:
        for idx, h in enumerate(headers):
            if kw in h:
                target_col = idx
                break
        if target_col is not None:
            break
    if target_col is None:
        raise ValueError(f"None of {keywords} found in headers: {headers}")
    values = []
    for row in ws.iter_rows(min_row=2, min_col=target_col + 1, max_col=target_col + 1):
        cell = row[0]
        if cell.value is not None:
            try:
                values.append(float(cell.value))
            except (ValueError, TypeError):
                pass
    wb.close()
    return values


def _read_csv_col(fname: str, keywords: list[str]) -> list[float]:
    """Read column from CSV where header contains any keyword."""
    with open(fname, encoding='utf-8-sig') as f:
        reader = csv.DictReader(f)
        headers = list(reader.fieldnames or [])
        target = None
        for kw in keywords:
            for h in headers:
                if kw in h:
                    target = h
                    break
            if target:
                break
        if target is None:
            raise ValueError(f"None of {keywords} found in headers: {headers}")
        values = []
        for row in reader:
            try:
                values.append(float(row[target]))
            except (ValueError, TypeError):
                pass
    return values


import os
import yaml


class OfflineEnv(gym.Env):
    """Gymnasium env running on local scenario data. Matches sandbox engine formulas."""

    DT_HOURS = 0.25       # 15-minute timestep
    POINTS_PER_DAY = 96

    def __init__(self, config_path: str, days: int = 7, max_episode_steps: int | None = None):
        # Load YAML config
        with open(config_path, 'r', encoding='utf-8') as f:
            yc = yaml.safe_load(f)

        sc = yc.get('scenario', {})
        dv = yc.get('device', {})
        bt = yc.get('battery', {})
        gd = yc.get('grid', {})
        cs = yc.get('cost', {})

        # Resolve paths relative to project root
        root = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
        pv_wind_file = os.path.join(root, sc['pv_wind_file'])
        load_file    = os.path.join(root, sc['load_file'])
        price_file   = os.path.join(root, sc['price_file'])

        # Load scenario data
        pv_raw  = _read_xlsx_col(pv_wind_file, ['辐照', 'irradiance', 'ghi', 'solar'])
        wind_raw = _read_xlsx_col(pv_wind_file, ['风速', 'wind'])
        load_raw = _read_xlsx_col(load_file,      ['负荷', 'load', '有功', 'demand'])
        buy_raw  = _read_csv_col(price_file,       ['购电价', 'buy_price'])
        sell_raw = _read_csv_col(price_file,       ['售电价', 'sell_price'])

        # Trim to days
        T_max = min(len(pv_raw), len(wind_raw), len(load_raw), len(buy_raw), len(sell_raw))
        T = min(days * self.POINTS_PER_DAY, T_max)

        # Convert units (matches sandbox instance.py:_apply_scenario_point)
        pv_cap = dv.get('pv_capacity_kw', 80.0)
        pv_eff = dv.get('pv_efficiency', 0.95)
        self.pv_kw = [min(1.0, p / 1000.0) * pv_cap * pv_eff for p in pv_raw[:T]]

        wind_cap = dv.get('wind_capacity_kw', 50.0)
        wind_eff = dv.get('wind_efficiency', 0.92)
        self.wind_kw = []
        for w_kmh in wind_raw[:T]:
            ws = w_kmh / 3.6           # km/h -> m/s
            ratio = min(1.0, ws / 12.0)
            self.wind_kw.append(ratio * wind_cap * wind_eff)

        load_base = dv.get('load_base_kw', 80.0)
        self.load_kw = [max(0.0, l * load_base) for l in load_raw[:T]]

        self.buy_price  = buy_raw[:T]
        self.sell_price = sell_raw[:T]

        # Battery params
        self.battery_capacity_kwh    = bt.get('capacity_kwh', 200.0)
        self.battery_charge_max_kw   = bt.get('charge_max_kw', 100.0)
        self.battery_discharge_max_kw = bt.get('discharge_max_kw', 100.0)
        self.battery_soc_min = bt.get('soc_min', 0.10)
        self.battery_soc_max = bt.get('soc_max', 0.95)
        self.battery_soc_init = bt.get('soc_init', 0.50)
        self.battery_charge_eff    = bt.get('charge_eff', 0.95)
        self.battery_discharge_eff = bt.get('discharge_eff', 0.95)

        # Grid params
        self.grid_import_max_kw = gd.get('import_max_kw', 100.0)
        self.grid_export_max_kw = gd.get('export_max_kw', 100.0)

        # Cost params
        self.c_deg = cs.get('c_deg', 0.05)

        # State
        self.T = T
        self.max_episode_steps = max_episode_steps or T
        self._t = 0
        self._soc = self.battery_soc_init
        self.obs_dim = 7
        self.action_space = spaces.Box(low=-1.0, high=1.0, shape=(1,), dtype=np.float32)
        self.observation_space = spaces.Box(
            low=-1.0, high=1000.0, shape=(self.obs_dim,), dtype=np.float32,
        )

        # Accumulators for evaluation
        self.total_purchase = 0.0
        self.total_revenue  = 0.0
        self.total_degradation = 0.0
        self.peak_grid_import = 0.0

    def reset(self, seed=None, options=None):
        super().reset(seed=seed)
        self._t = 0
        self._soc = self.battery_soc_init
        self.total_purchase = 0.0
        self.total_revenue  = 0.0
        self.total_degradation = 0.0
        self.peak_grid_import = 0.0
        return self._get_obs(), {}

    def step(self, action):
        # Handle various action formats (list, tuple, numpy array)
        if isinstance(action, np.ndarray):
            action = float(action.flat[0])
        elif isinstance(action, (list, tuple)):
            action = float(action[0])
        else:
            action = float(action)

        # Clip action to battery limits
        battery_power = max(-self.battery_charge_max_kw,
                           min(self.battery_discharge_max_kw, action))

        # Power balance: PV + Wind + Bat + Grid = Load
        grid = (self.load_kw[self._t] - self.pv_kw[self._t]
                - self.wind_kw[self._t] - battery_power)
        grid_in  = max(0.0,  grid)
        grid_out = max(0.0, -grid)

        # SOC dynamics (matches sandbox engine.py eq_soc_dynamics)
        chg = max(0.0, -battery_power)
        dis = max(0.0,  battery_power)
        delta_soc = (self.battery_charge_eff * chg
                     - dis / self.battery_discharge_eff) * self.DT_HOURS / self.battery_capacity_kwh
        new_soc = self._soc + delta_soc

        # Violation penalty: if SOC goes out of bounds, clip and penalize
        soc_violation = 0.0
        if new_soc < self.battery_soc_min:
            soc_violation = (self.battery_soc_min - new_soc) * self.battery_capacity_kwh * 100.0
            new_soc = self.battery_soc_min
        elif new_soc > self.battery_soc_max:
            soc_violation = (new_soc - self.battery_soc_max) * self.battery_capacity_kwh * 100.0
            new_soc = self.battery_soc_max
        self._soc = new_soc

        # Reward = negative of MILP objective (one step)
        dt = self.DT_HOURS
        purchase  = self.buy_price[self._t] * grid_in * dt
        revenue   = self.sell_price[self._t] * grid_out * dt
        degradation = self.c_deg * (chg + dis) * dt
        reward = -(purchase - revenue + degradation + soc_violation)

        # Accumulate for evaluation
        self.total_purchase    += purchase
        self.total_revenue     += revenue
        self.total_degradation += degradation
        if grid_in > self.peak_grid_import:
            self.peak_grid_import = grid_in

        self._t += 1
        terminated = self._t >= self.T
        truncated = self._t >= self.max_episode_steps

        return self._get_obs(), reward, terminated, truncated, {'soc_violation': soc_violation > 0}

    def _get_obs(self):
        idx = min(self._t, self.T - 1)  # clamp to last valid timestep
        return [
            self._soc,
            self.pv_kw[idx] / 100.0,      # normalize ~0-1
            self.wind_kw[idx] / 100.0,
            self.load_kw[idx] / 100.0,
            self.buy_price[idx],
            self.sell_price[idx],
            (self._t % self.POINTS_PER_DAY) / self.POINTS_PER_DAY,  # time-of-day 0~1
        ]

    @property
    def total_cost(self) -> float:
        """Total cost for the episode (matching MILP comprehensive cost format)."""
        return self.total_purchase - self.total_revenue + self.total_degradation
